# -*- coding: utf-8 -*-
"""
阶段5：父子文档切分 + MRR 对比评估

要回答的问题：父子文档切分（子块200字符召回 → 父块800字符回答）相比
单层固定切分（400字符），MRR 提升多少？

评估设计（为什么这么做）：
- ground-truth：答案数值在语料里「唯一」（命中 1~3 块）的干净题，共约17条。
  这些题答案是具体金额（如 318673.34），用「答案数值包含」做相关性判据是可靠的。
- 封闭集：只用包含这些答案的文档（约23个），两种切分在【同一文档集】上比，
  保证变量只有「切分方式」一个。
- 检索链两种切分完全一致：混合检索(BM25+向量 RRF) → cross-encoder 重排。
- MRR = 平均 1/首个含答案的块的排名（父子：子块召回后映射回父块判相关性）。

用法：.venv\\Scripts\\python.exe stage5_eval_parent_child.py
"""
import os
os.environ.setdefault("HF_HOME", "E:/hf_cache")
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("NO_PROXY", "*")
os.environ.setdefault("no_proxy", "*")

import json, re
from pathlib import Path
import pymupdf, faiss, numpy as np, jieba
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer, CrossEncoder

BASE = Path("E:/数字员工项目")
DOCS = BASE / "docs"
CHUNKS_FILE = BASE / "chunks.json"
EVAL_FILE = BASE / "评估集.json"
K = 10  # MRR 的 top-k 截断

print("加载 embedding + reranker 模型 ...", flush=True)
embed_model = SentenceTransformer("BAAI/bge-small-zh-v1.5")
reranker = CrossEncoder("BAAI/bge-reranker-base")
print("模型就绪", flush=True)

# ================= 文档解析（与 rebuild_index.py 一致） =================
def extract_lines(page, y_tol=3.0):
    def is_data(w):
        return bool(re.match(r'^[\d±—✓\-\.\(\)]', w))
    words = page.get_text("words")
    words.sort(key=lambda w: (round(w[1], 1), w[0]))
    raw_lines, cur, cur_y = [], [], None
    for w in words:
        y = w[1]
        if cur_y is None or abs(y - cur_y) > y_tol:
            if cur:
                raw_lines.append(cur)
            cur, cur_y = [w[4]], y
        else:
            cur.append(w[4])
    if cur:
        raw_lines.append(cur)
    result = []
    for ws in raw_lines:
        dc = sum(1 for w in ws if is_data(w))
        if dc >= 2 and len(ws) - dc >= 1:
            labels = [w for w in ws if not is_data(w)]
            datas = [w for w in ws if is_data(w)]
            result.append(" ".join(labels + datas))
        else:
            result.append(" ".join(ws))
    return result

def load_pdf(path):
    doc = pymupdf.open(str(path))
    lines = []
    for page in doc:
        lines.extend(extract_lines(page))
    return lines

def load_document(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(str(path))
    elif ext in [".txt", ".md"]:
        return Path(path).read_text(encoding="utf-8", errors="replace").splitlines()
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [para.text for para in d.paragraphs if para.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(cells))
        return lines
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[工作表: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        wb.close()
        return lines
    return []

# ================= 切分：单层 vs 父子 =================
def split_flat(lines, chunk_size=400):
    """现状：单层固定大小切块（按行累积，行内不切断）"""
    chunks, current, cur_len = [], [], 0
    for line in lines:
        if current and cur_len + len(line) > chunk_size:
            chunks.append("\\n".join(current))
            current, cur_len = [], 0
        current.append(line)
        cur_len += len(line)
    if current:
        chunks.append("\\n".join(current))
    return chunks

def split_parent_child(lines, child_size=200, parent_size=800):
    """父子切块：先切父块(800)，每个父块内再切子块(200)。
    返回 (子块列表, 子块->父块下标映射, 父块列表)"""
    # 1) 先按 parent_size 切父块
    parents, cur, cur_len = [], [], 0
    for line in lines:
        if cur and cur_len + len(line) > parent_size:
            parents.append("\\n".join(cur))
            cur, cur_len = [], 0
        cur.append(line)
        cur_len += len(line)
    if cur:
        parents.append("\\n".join(cur))
    # 2) 每个父块内按 child_size 切子块
    children, child2parent = [], []
    for pi, ptext in enumerate(parents):
        cc, cclen = [], 0
        for ln in ptext.split("\\n"):
            if cc and cclen + len(ln) > child_size:
                children.append("\\n".join(cc))
                child2parent.append(pi)
                cc, cclen = [], 0
            cc.append(ln)
            cclen += len(ln)
        if cc:
            children.append("\\n".join(cc))
            child2parent.append(pi)
    return children, child2parent, parents

# ================= 索引 + 检索 =================
def build_faiss(texts):
    vectors = np.array(embed_model.encode(texts)).astype("float32")
    idx = faiss.IndexFlatL2(vectors.shape[1])
    idx.add(vectors)
    return idx

def build_bm25(texts):
    return BM25Okapi([list(jieba.cut(t)) for t in texts])

def hybrid_search(query, faiss_idx, bm25, n=20):
    q_vec = embed_model.encode([query]).astype("float32")
    _, vec_ids = faiss_idx.search(q_vec, n)
    vec_ids = vec_ids[0].tolist()
    bm25_ids = np.argsort(bm25.get_scores(list(jieba.cut(query))))[::-1][:n].tolist()
    rrf = {}
    for rank, i in enumerate(vec_ids):
        rrf[i] = rrf.get(i, 0) + 1.0 / (60 + rank)
    for rank, i in enumerate(bm25_ids):
        rrf[i] = rrf.get(i, 0) + 1.0 / (60 + rank)
    return sorted(rrf, key=rrf.get, reverse=True)

def rerank_top(query, candidates, texts, k):
    pairs = [(query, texts[i]) for i in candidates]
    scores = reranker.predict(pairs)
    ranked = sorted(zip(candidates, scores), key=lambda x: x[1], reverse=True)
    return [i for i, s in ranked[:k]]

# ================= 归一化 + 答案数值 =================
def norm(s):
    return (s.replace(" ", "").replace("\u3000", "").replace("\\n", "")
             .replace("\\r", "").replace("\\t", "")
             .replace(",", "").replace("，", ""))

def answer_value(ans):
    """取答案里最具体的数字（优先带小数、位数多的），用于相关性判据"""
    nums = re.findall(r'\d+(?:\.\d+)?', ans)
    if not nums:
        return None
    return max(nums, key=lambda n: (len(n), '.' in n))

# ================= 1. 用全量语料找干净题 + 范围文档 =================
evals = json.loads(EVAL_FILE.read_text(encoding="utf-8"))
full = json.loads(CHUNKS_FILE.read_text(encoding="utf-8"))
full_texts = [c["text"] for c in full]
full_sources = [c["source"] for c in full]
full_norm = [norm(t) for t in full_texts]

clean = []  # (eval_item, value, golden_full_indices)
for it in evals:
    if "无此信息" in it["a"] or it["a"].startswith("文档未"):
        continue
    v = answer_value(it["a"])
    if v is None:
        continue
    nv = norm(v)
    if len(nv) < 4 and "." not in v:   # 太短无区分度
        continue
    idx = [i for i, nt in enumerate(full_norm) if nv in nt]
    if 1 <= len(idx) <= 3:              # 唯一性判据
        clean.append((it, v, idx))

scoped_docs = sorted(set(full_sources[i] for _, _, idx in clean for i in idx))
print(f"干净题: {len(clean)} 条；范围文档: {len(scoped_docs)} 个", flush=True)

# ================= 2. 解析范围文档，两种切分 =================
flat_chunks, flat_sources = [], []
child_chunks, child2parent, parents, child_sources = [], [], [], []

for rel in scoped_docs:
    path = DOCS / rel
    try:
        lines = load_document(str(path))
    except Exception as e:
        print(f"  解析失败 {rel}: {type(e).__name__}", flush=True)
        continue
    # 单层
    fc = split_flat(lines)
    flat_chunks.extend(fc)
    flat_sources.extend([rel] * len(fc))
    # 父子
    cc, c2p, pp = split_parent_child(lines)
    base = len(parents)
    child_chunks.extend(cc)
    child2parent.extend([base + p for p in c2p])
    parents.extend(pp)
    child_sources.extend([rel] * len(cc))

print(f"单层: {len(flat_chunks)} 块；父子: 父 {len(parents)} 块 / 子 {len(child_chunks)} 块", flush=True)

# ================= 3. 建两种索引 =================
print("编码单层块 ...", flush=True)
fa_flat = build_faiss(flat_chunks)
bm_flat = build_bm25(flat_chunks)
print("编码父子子块 ...", flush=True)
fa_child = build_faiss(child_chunks)
bm_child = build_bm25(child_chunks)

flat_norm = [norm(t) for t in flat_chunks]
parent_norm = [norm(t) for t in parents]

# ================= 4. 检索 + MRR =================
def reciprocal_rank(q, value, fa, bm, texts, norm_texts, is_parent_child=False):
    """返回该题的 reciprocal rank（0=top-K 没找到）"""
    nv = norm(value)
    cand = hybrid_search(q, fa, bm)[:20]
    top = rerank_top(q, cand, texts, K)
    if is_parent_child:
        # 子块 -> 映射父块（去重，保序）
        seen, ranked_parents = set(), []
        for ci in top:
            pi = child2parent[ci]
            if pi not in seen:
                seen.add(pi)
                ranked_parents.append(pi)
        for rank, pi in enumerate(ranked_parents, 1):
            if nv in parent_norm[pi]:
                return 1.0 / rank
    else:
        for rank, i in enumerate(top, 1):
            if nv in norm_texts[i]:
                return 1.0 / rank
    return 0.0

print(f"\n{'id':>3} {'类别':<3} {'答案值':<12} {'单层RR':>7} {'父子RR':>7}", flush=True)
flat_rrs, pc_rrs = [], []
for it, v, _ in clean:
    q = it["q"]
    rr_flat = reciprocal_rank(q, v, fa_flat, bm_flat, flat_chunks, flat_norm, False)
    rr_pc = reciprocal_rank(q, v, fa_child, bm_child, child_chunks, None, True)
    flat_rrs.append(rr_flat)
    pc_rrs.append(rr_pc)
    flag = "▲" if rr_pc > rr_flat else ("▼" if rr_pc < rr_flat else "=")
    print(f"{it['id']:>3} {it['cat']:<3} {v:<12} {rr_flat:>7.4f} {rr_pc:>7.4f} {flag}", flush=True)

mrr_flat = sum(flat_rrs) / len(flat_rrs)
mrr_pc = sum(pc_rrs) / len(pc_rrs)
print(f"\n===== 结果 =====", flush=True)
print(f"MRR@K(K={K})  单层={mrr_flat:.4f}  父子={mrr_pc:.4f}", flush=True)
print(f"提升: {mrr_pc - mrr_flat:+.4f}  ({(mrr_pc - mrr_flat) / mrr_flat * 100:+.1f}%)", flush=True)
print(f"找到答案的题数: 单层 {sum(1 for r in flat_rrs if r > 0)}/{len(flat_rrs)}  父子 {sum(1 for r in pc_rrs if r > 0)}/{len(pc_rrs)}", flush=True)
