# -*- coding: utf-8 -*-
"""
阶段11：混合测试集（50%规章制度+50%其他），「全单层」vs「二元路由」MRR
- 方法A 全单层400
- 方法B 二元路由：散文(占比>0.7)走父子200/800，表格走单层400
用法：.venv\\Scripts\\python.exe stage11_mixed.py
"""
import os
os.environ.setdefault("HF_HOME", "E:/hf_cache")
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("NO_PROXY", "*")
os.environ.setdefault("no_proxy", "*")

import json, re
from pathlib import Path
import pymupdf, faiss, numpy as np, jieba
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer, CrossEncoder

BASE = Path("E:/数字员工项目")
DOCS = BASE / "docs"
EVAL = BASE / "评估集_混合.json"
K = 10
TH = 0.7

print("加载模型 ...", flush=True)
embed_model = SentenceTransformer("BAAI/bge-small-zh-v1.5")
reranker = CrossEncoder("BAAI/bge-reranker-base")
print("模型就绪", flush=True)

def extract_lines(page, y_tol=3.0):
    def is_data(w): return bool(re.match(r'^[\d±—✓\-\.\(\)]', w))
    words = page.get_text("words"); words.sort(key=lambda w:(round(w[1],1), w[0]))
    raw, cur, cy = [], [], None
    for w in words:
        y = w[1]
        if cy is None or abs(y-cy) > y_tol:
            if cur: raw.append(cur)
            cur, cy = [w[4]], y
        else: cur.append(w[4])
    if cur: raw.append(cur)
    out = []
    for ws in raw:
        dc = sum(1 for w in ws if is_data(w))
        out.append(" ".join([w for w in ws if not is_data(w)] + [w for w in ws if is_data(w)]) if dc>=2 and len(ws)-dc>=1 else " ".join(ws))
    return out

def load_document(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        doc = pymupdf.open(str(path)); lines = []
        for pg in doc: lines.extend(extract_lines(pg))
        return lines
    if ext in (".txt",".md"): return Path(path).read_text(encoding="utf-8", errors="replace").splitlines()
    if ext == ".docx":
        import docx
        d = docx.Document(str(path)); lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows: lines.append(" | ".join(c.text.strip() for c in r.cells))
        return lines
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True); lines = []
        for ws in wb.worksheets:
            lines.append(f"[工作表: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells): lines.append(" | ".join(cells))
        wb.close(); return lines
    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path); lines = []
        for sh in wb.sheets():
            lines.append(f"[工作表: {sh.name}]")
            for r in range(sh.nrows):
                cells = [str(sh.cell_value(r,c)) for c in range(sh.ncols)]
                if any(c.strip() for c in cells): lines.append(" | ".join(cells))
        return lines
    return []

def prose_ratio(path):
    ext = Path(path).suffix.lower()
    if ext == ".xlsx": return 0.0
    if ext in (".txt",".md"): return 1.0
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        p = sum(len(x.text) for x in d.paragraphs)
        t = sum(len(c.text) for tb in d.tables for r in tb.rows for c in r.cells)
        return p/(p+t) if (p+t) else 0.5
    return 0.5

def split_flat(lines, size=400):
    chunks, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen+len(ln) > size: chunks.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: chunks.append("\n".join(cur))
    return chunks

def split_pc(lines, child=200, parent=800):
    parents, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen+len(ln) > parent: parents.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: parents.append("\n".join(cur))
    children, c2p = [], []
    for pi, p in enumerate(parents):
        cc, cclen = [], 0
        for ln in p.split("\n"):
            if cc and cclen+len(ln) > child: children.append("\n".join(cc)); c2p.append(pi); cc, cclen = [], 0
            cc.append(ln); cclen += len(ln)
        if cc: children.append("\n".join(cc)); c2p.append(pi)
    return children, c2p, parents

def build_faiss(texts):
    v = np.array(embed_model.encode(texts)).astype("float32")
    idx = faiss.IndexFlatL2(v.shape[1]); idx.add(v); return idx
def build_bm25(texts): return BM25Okapi([list(jieba.cut(t)) for t in texts])
def hybrid(q, fa, bm, n=20):
    qv = embed_model.encode([q]).astype("float32"); _, vid = fa.search(qv, n); vid = vid[0].tolist()
    bid = np.argsort(bm.get_scores(list(jieba.cut(q))))[::-1][:n].tolist()
    rrf = {}
    for r,i in enumerate(vid): rrf[i] = rrf.get(i,0)+1/(60+r)
    for r,i in enumerate(bid): rrf[i] = rrf.get(i,0)+1/(60+r)
    return sorted(rrf, key=rrf.get, reverse=True)
def rerank(q, cands, texts, k):
    sc = reranker.predict([(q, texts[i]) for i in cands])
    return [i for i,s in sorted(zip(cands, sc), key=lambda x:x[1], reverse=True)[:k]]
def norm(s):
    return (s.replace(" ","").replace("\u3000","").replace("\n","").replace("\r","").replace("\t","").replace(",","").replace("，",""))

# 语料
files = [f for f in sorted(DOCS.rglob("*")) if f.is_file() and not f.name.startswith("~$")
         and f.suffix.lower() in (".docx",".xlsx",".xls",".pdf")
         and (("员工手册" in f.name) or ("施工组织方案" in f.name) or ("专项施工方案" in f.name)
              or f.suffix.lower() in (".xlsx",".xls") or ("运行部零星工程.PDF" in f.name))]
print(f"语料 {len(files)} 文件", flush=True)

flat_A = []          # 方法A：全单层块
units_B = []         # 方法B：检索单元（子块或平块）
ctx_B = []           # 方法B：每个单元对应的"回答上下文"（父块或平块自身）

for f in files:
    try:
        lines = load_document(str(f))
    except Exception:
        continue
    if not lines: continue
    flat_A.extend(split_flat(lines))
    if prose_ratio(str(f)) > TH:
        ch, c2p, par = split_pc(lines)
        for i, c in enumerate(ch):
            units_B.append(c)
            ctx_B.append(par[c2p[i]])
    else:
        for c in split_flat(lines):
            units_B.append(c); ctx_B.append(c)

print(f"方法A 单层 {len(flat_A)} 块 | 方法B 检索单元 {len(units_B)} (上下文 {len(ctx_B)})", flush=True)
print("编码 ...", flush=True)
faA = build_faiss(flat_A); bmA = build_bm25(flat_A)
faB = build_faiss(units_B); bmB = build_bm25(units_B)
normA = [norm(t) for t in flat_A]
normCtxB = [norm(t) for t in ctx_B]

evals = json.loads(EVAL.read_text(encoding="utf-8"))
# 唯一性过滤（基于方法A语料）
clean = [(it, norm(it["a"])) for it in evals
         if 1 <= sum(1 for t in normA if norm(it["a"]) in t) <= 8]
print(f"可评估 {len(clean)}/{len(evals)} 题", flush=True)

def rrA(q, nv):
    for rank, i in enumerate(rerank(q, hybrid(q, faA, bmA)[:20], flat_A, K), 1):
        if nv in normA[i]: return 1.0/rank
    return 0.0
def rrB(q, nv):
    for rank, i in enumerate(rerank(q, hybrid(q, faB, bmB)[:20], units_B, K), 1):
        if nv in normCtxB[i]: return 1.0/rank
    return 0.0

from collections import defaultdict
cat_stat = defaultdict(lambda: [0.0,0.0,0])  # cat -> [sumA, sumB, n]
print(f"\n{'答案':<22} {'类':<6} {'全单层':>7} {'二元路由':>7}", flush=True)
for it, nv in clean:
    a = rrA(it["q"], nv); b = rrB(it["q"], nv)
    c = it.get("cat","?")
    cat_stat[c][0]+=a; cat_stat[c][1]+=b; cat_stat[c][2]+=1
    print(f"{it['a']:<22} {c:<6} {a:>7.3f} {b:>7.3f}", flush=True)

mA = sum(x[0] for x in cat_stat.values())/sum(x[2] for x in cat_stat.values())
mB = sum(x[1] for x in cat_stat.values())/sum(x[2] for x in cat_stat.values())
print(f"\n===== 混合 MRR@10 =====", flush=True)
print(f"全单层   = {mA:.4f}", flush=True)
print(f"二元路由 = {mB:.4f}  ({mB-mA:+.4f})", flush=True)
print("\n分类型：", flush=True)
for c, (sa, sb, n) in sorted(cat_stat.items()):
    print(f"  {c:<8} 全单层 {sa/n:.4f}  二元路由 {sb/n:.4f}  ({(sb-sa)/n:+.4f})  n={n}", flush=True)
