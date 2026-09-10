# -*- coding: utf-8 -*-
"""
阶段10：长文本（施工方案）单层 vs 父子，双判据 MRR
① chunk级：答案在语料命中≤5块的题（唯一性较松）
② 文档级：检索块来源 == 标注的 ground-truth source
用法：.venv\\Scripts\\python.exe stage10_longform2.py
"""
import os
os.environ.setdefault("HF_HOME", "E:/hf_cache")
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("NO_PROXY", "*")
os.environ.setdefault("no_proxy", "*")

import json, re
from pathlib import Path
import faiss, numpy as np, jieba
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer, CrossEncoder

BASE = Path("E:/数字员工项目")
DOCS = BASE / "docs"
EVAL = BASE / "评估集_长文本.json"
K = 10

print("加载模型 ...", flush=True)
embed_model = SentenceTransformer("BAAI/bge-small-zh-v1.5")
reranker = CrossEncoder("BAAI/bge-reranker-base")
print("模型就绪", flush=True)

def load_docx(path):
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return lines

def split_flat(lines, size=400):
    chunks, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen + len(ln) > size:
            chunks.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: chunks.append("\n".join(cur))
    return chunks

def split_parent_child(lines, child=200, parent=800):
    parents, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen + len(ln) > parent:
            parents.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: parents.append("\n".join(cur))
    children, c2p = [], []
    for pi, p in enumerate(parents):
        cc, cclen = [], 0
        for ln in p.split("\n"):
            if cc and cclen + len(ln) > child:
                children.append("\n".join(cc)); c2p.append(pi); cc, cclen = [], 0
            cc.append(ln); cclen += len(ln)
        if cc: children.append("\n".join(cc)); c2p.append(pi)
    return children, c2p, parents

def build_faiss(texts):
    v = np.array(embed_model.encode(texts)).astype("float32")
    idx = faiss.IndexFlatL2(v.shape[1]); idx.add(v); return idx
def build_bm25(texts):
    return BM25Okapi([list(jieba.cut(t)) for t in texts])
def hybrid_search(q, fa, bm, n=20):
    qv = embed_model.encode([q]).astype("float32"); _, vid = fa.search(qv, n); vid = vid[0].tolist()
    bid = np.argsort(bm.get_scores(list(jieba.cut(q))))[::-1][:n].tolist()
    rrf = {}
    for r,i in enumerate(vid): rrf[i] = rrf.get(i,0) + 1/(60+r)
    for r,i in enumerate(bid): rrf[i] = rrf.get(i,0) + 1/(60+r)
    return sorted(rrf, key=rrf.get, reverse=True)
def rerank_top(q, cands, texts, k):
    sc = reranker.predict([(q, texts[i]) for i in cands])
    return [i for i,s in sorted(zip(cands, sc), key=lambda x:x[1], reverse=True)[:k]]
def norm(s):
    return (s.replace(" ","").replace("\u3000","").replace("\n","").replace("\r","")
             .replace("\t","").replace(",","").replace("，",""))

# 语料
files = [f for f in sorted(DOCS.rglob("*.docx"))
         if ("施工组织方案" in f.name or "专项施工方案" in f.name) and not f.name.startswith("~$")]
print(f"语料 {len(files)} 文件", flush=True)

flat, flat_src, pc_ch, pc_c2p, pc_par, pc_par_src = [], [], [], [], [], []
for f in files:
    rel = str(f.relative_to(DOCS))
    try:
        lines = load_docx(f)
    except Exception:
        continue
    fc = split_flat(lines)
    flat.extend(fc); flat_src.extend([rel]*len(fc))
    c, c2p, p = split_parent_child(lines)
    base = len(pc_par)
    pc_ch.extend(c); pc_c2p.extend([base+x for x in c2p]); pc_par.extend(p)
    pc_par_src.extend([rel]*len(p))

print(f"单层 {len(flat)} 块 | 父子 子{len(pc_ch)}/父{len(pc_par)}", flush=True)
print("编码 ...", flush=True)
fa_f = build_faiss(flat); bm_f = build_bm25(flat)
fa_c = build_faiss(pc_ch); bm_c = build_bm25(pc_ch)
flat_norm = [norm(t) for t in flat]
pc_par_norm = [norm(t) for t in pc_par]

evals = json.loads(EVAL.read_text(encoding="utf-8"))

def top_flat(q):
    cand = hybrid_search(q, fa_f, bm_f)[:20]
    return rerank_top(q, cand, flat, K)
def top_pc_parents(q):
    cand = hybrid_search(q, fa_c, bm_c)[:20]
    seen, rp = set(), []
    for ci in rerank_top(q, cand, pc_ch, K):
        pi = pc_c2p[ci]
        if pi not in seen:
            seen.add(pi); rp.append(pi)
    return rp

# ---- ① chunk级（答案命中≤5块） ----
print(f"\n===== ① chunk级 MRR（答案≤5块） =====", flush=True)
cand_evals = [(it, norm(it["a"])) for it in evals
              if 1 <= sum(1 for nt in flat_norm if norm(it["a"]) in nt) <= 5]
print(f"可评估 {len(cand_evals)}/{len(evals)} 题\n", flush=True)
r1, r2 = [], []
for it, nv in cand_evals:
    a = 0.0; b = 0.0
    for rank, i in enumerate(top_flat(it["q"]), 1):
        if nv in flat_norm[i]: a = 1.0/rank; break
    for rank, pi in enumerate(top_pc_parents(it["q"]), 1):
        if nv in pc_par_norm[pi]: b = 1.0/rank; break
    r1.append(a); r2.append(b)
    print(f"  {it['a']:<8} 单层{a:.3f} 父子{b:.3f}", flush=True)
m1 = sum(r1)/len(r1); m2 = sum(r2)/len(r2)
print(f"  chunk级 MRR@10: 单层={m1:.4f}  父子={m2:.4f}  ({m2-m1:+.4f})", flush=True)

# ---- ② 文档级（source 匹配） ----
print(f"\n===== ② 文档级 MRR（source 匹配） =====", flush=True)
d1, d2 = [], []
for it in evals:
    gt = it["source"]
    a = 0.0; b = 0.0
    for rank, i in enumerate(top_flat(it["q"]), 1):
        if flat_src[i] == gt: a = 1.0/rank; break
    for rank, pi in enumerate(top_pc_parents(it["q"]), 1):
        if pc_par_src[pi] == gt: b = 1.0/rank; break
    d1.append(a); d2.append(b)
dm1 = sum(d1)/len(d1); dm2 = sum(d2)/len(d2)
print(f"  文档级 MRR@10: 单层={dm1:.4f}  父子={dm2:.4f}  ({dm2-dm1:+.4f})", flush=True)
print(f"  文档级命中: 单层 {sum(1 for x in d1 if x>0)}/{len(d1)}  父子 {sum(1 for x in d2 if x>0)}/{len(d2)}", flush=True)
