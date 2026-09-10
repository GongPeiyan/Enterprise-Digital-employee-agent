# -*- coding: utf-8 -*-
"""
阶段7：按行切 + 工程名上下文（contextual chunking）对比
修正阶段6的问题：naive 按行切把「工程名」和「子项金额」拆到不同块，丢了上下文。
这里按行切的同时，把工程名拼回每一行，再比 MRR。

用法：.venv\\Scripts\\python.exe stage7_context_row.py
"""
import os
os.environ.setdefault("HF_HOME", "E:/hf_cache")
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("NO_PROXY", "*")
os.environ.setdefault("no_proxy", "*")

import json, re
from pathlib import Path
import pymupdf, faiss, numpy as np, jieba
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer, CrossEncoder

BASE = Path("E:/数字员工项目")
DOCS = BASE / "docs"
CHUNKS_FILE = BASE / "chunks.json"
EVAL_FILE = BASE / "评估集.json"
SUPPORTED = {".pdf", ".txt", ".md", ".docx", ".xlsx"}
K = 10

print("加载模型 ...", flush=True)
embed_model = SentenceTransformer("BAAI/bge-small-zh-v1.5")
reranker = CrossEncoder("BAAI/bge-reranker-base")
print("模型就绪", flush=True)

def extract_lines(page, y_tol=3.0):
    def is_data(w):
        return bool(re.match(r'^[\d±—✓\-\.\(\)]', w))
    words = page.get_text("words"); words.sort(key=lambda w:(round(w[1],1), w[0]))
    raw, cur, cy = [], [], None
    for w in words:
        y = w[1]
        if cy is None or abs(y-cy) > y_tol:
            if cur: raw.append(cur)
            cur, cy = [w[4]], y
        else:
            cur.append(w[4])
    if cur: raw.append(cur)
    out = []
    for ws in raw:
        dc = sum(1 for w in ws if is_data(w))
        if dc >= 2 and len(ws)-dc >= 1:
            out.append(" ".join([w for w in ws if not is_data(w)] + [w for w in ws if is_data(w)]))
        else:
            out.append(" ".join(ws))
    return out

def load_pdf(path):
    doc = pymupdf.open(str(path)); lines = []
    for page in doc: lines.extend(extract_lines(page))
    return lines

def load_document(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf": return load_pdf(str(path))
    if ext in [".txt", ".md"]: return Path(path).read_text(encoding="utf-8", errors="replace").splitlines()
    if ext == ".docx":
        import docx
        d = docx.Document(str(path)); lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return lines
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True); lines = []
        for ws in wb.worksheets:
            lines.append(f"[工作表: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells): lines.append(" | ".join(cells))
        wb.close(); return lines
    return []

def split_flat(lines, size=400):
    chunks, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen + len(ln) > size:
            chunks.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: chunks.append("\n".join(cur))
    return chunks

def find_project(lines):
    """从行里找工程名（工程名称：/项目名称：后的内容）"""
    for ln in lines:
        m = re.search(r'(?:工程名称|项目名称)\s*[:：]\s*([^|｜]+)', ln)
        if m and m.group(1).strip():
            return m.group(1).strip()
    return None

def chunk_rows_context(lines):
    """按行切 + 工程名上下文拼回每行"""
    proj = find_project(lines)
    chunks = []
    for ln in lines:
        s = ln.strip()
        if not s or s.startswith("[工作表:"):
            continue
        chunks.append(f"{proj} | {s}" if proj else s)
    return chunks

def build_faiss(texts):
    v = np.array(embed_model.encode(texts)).astype("float32")
    idx = faiss.IndexFlatL2(v.shape[1]); idx.add(v); return idx
def build_bm25(texts):
    return BM25Okapi([list(jieba.cut(t)) for t in texts])
def hybrid_search(q, fa, bm, n=20):
    qv = embed_model.encode([q]).astype("float32"); _, vid = fa.search(qv, n); vid = vid[0].tolist()
    bid = np.argsort(bm.get_scores(list(jieba.cut(q))))[::-1][:n].tolist()
    rrf = {}
    for r,i in enumerate(vid): rrf[i] = rrf.get(i,0) + 1/(60+r)
    for r,i in enumerate(bid): rrf[i] = rrf.get(i,0) + 1/(60+r)
    return sorted(rrf, key=rrf.get, reverse=True)
def rerank_top(q, cands, texts, k):
    sc = reranker.predict([(q, texts[i]) for i in cands])
    return [i for i,s in sorted(zip(cands, sc), key=lambda x:x[1], reverse=True)[:k]]
def norm(s):
    return (s.replace(" ","").replace("\u3000","").replace("\n","").replace("\r","")
             .replace("\t","").replace(",","").replace("，",""))
def answer_value(ans):
    nums = re.findall(r'\d+(?:\.\d+)?', ans)
    if not nums: return None
    return max(nums, key=lambda n:(len(n), '.' in n))

evals = json.loads(EVAL_FILE.read_text(encoding="utf-8"))
full = json.loads(CHUNKS_FILE.read_text(encoding="utf-8"))
full_norm = [norm(c["text"]) for c in full]
full_src = [c["source"] for c in full]

clean = []
for it in evals:
    if "无此信息" in it["a"] or it["a"].startswith("文档未"): continue
    v = answer_value(it["a"])
    if v is None: continue
    nv = norm(v)
    if len(nv) < 4 and "." not in v: continue
    idx = [i for i,nt in enumerate(full_norm) if nv in nt]
    if 1 <= len(idx) <= 3: clean.append((it, v))

golden_docs = set()
for it, v in clean:
    nv = norm(v)
    for i, nt in enumerate(full_norm):
        if nv in nt: golden_docs.add(full_src[i])
golden_folders = set(Path(d).parent for d in golden_docs)
corpus_files = [f for f in sorted(DOCS.rglob("*")) if f.is_file() and not f.name.startswith("~$")
                and f.suffix.lower() in SUPPORTED and f.relative_to(DOCS).parent in golden_folders]

flat_chunks, ctxrow_chunks = [], []
for f in corpus_files:
    try:
        lines = load_document(str(f))
    except Exception:
        continue
    flat_chunks.extend(split_flat(lines))
    if f.suffix.lower() in (".xlsx", ".pdf"):
        ctxrow_chunks.extend(chunk_rows_context(lines))
    else:
        ctxrow_chunks.extend(split_flat(lines))  # 长文档保持单层

print(f"干净题 {len(clean)}；语料 {len(corpus_files)} 文件", flush=True)
print(f"单层 {len(flat_chunks)} 块 | 按行+上下文 {len(ctxrow_chunks)} 块", flush=True)
print("编码 ...", flush=True)
fa_f = build_faiss(flat_chunks); bm_f = build_bm25(flat_chunks)
fa_r = build_faiss(ctxrow_chunks); bm_r = build_bm25(ctxrow_chunks)
flat_norm = [norm(t) for t in flat_chunks]
ctxrow_norm = [norm(t) for t in ctxrow_chunks]

def rr_flat(q, v):
    nv = norm(v); cand = hybrid_search(q, fa_f, bm_f)[:20]
    for rank, i in enumerate(rerank_top(q, cand, flat_chunks, K), 1):
        if nv in flat_norm[i]: return 1.0/rank
    return 0.0
def rr_ctxrow(q, v):
    nv = norm(v); cand = hybrid_search(q, fa_r, bm_r)[:20]
    for rank, i in enumerate(rerank_top(q, cand, ctxrow_chunks, K), 1):
        if nv in ctxrow_norm[i]: return 1.0/rank
    return 0.0

print(f"\n{'id':>3} {'类别':<3} {'答案值':<12} {'单层RR':>7} {'按行+上下文RR':>7}", flush=True)
r1, r2 = [], []
for it, v in clean:
    a = rr_flat(it["q"], v); b = rr_ctxrow(it["q"], v)
    r1.append(a); r2.append(b)
    print(f"{it['id']:>3} {it['cat']:<3} {v:<12} {a:>7.4f} {b:>7.4f}", flush=True)

m1 = sum(r1)/len(r1); m2 = sum(r2)/len(r2)
print(f"\n===== MRR@10 =====", flush=True)
print(f"单层400      = {m1:.4f}", flush=True)
print(f"按行+工程名上下文 = {m2:.4f}  ({m2-m1:+.4f})", flush=True)
print(f"命中: 单层 {sum(1 for x in r1 if x>0)}/{len(r1)}  按行+上下文 {sum(1 for x in r2 if x>0)}/{len(r2)}", flush=True)
