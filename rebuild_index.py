# -*- coding: utf-8 -*-
"""独立重建向量库脚本：递归扫描 docs/ 下所有支持格式，逐文件入库并打印进度。
用法：.venv\Scripts\python.exe rebuild_index.py
"""
import os
os.environ.setdefault("HF_HOME", "E:/hf_cache")
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("HF_HUB_OFFLINE", "1")
os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")
os.environ.setdefault("NO_PROXY", "*")
os.environ.setdefault("no_proxy", "*")

import json, re, sys
from pathlib import Path
import pymupdf, faiss, numpy as np, jieba
from sentence_transformers import SentenceTransformer

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))
import index_meta

DOCS_DIR = index_meta.docs_dir()          # 默认 BASE/docs，KB_DOCS_DIR 可覆盖（测试用）
INDEX_FILE = index_meta.index_file()      # KB_INDEX_FILE 可覆盖
CHUNKS_FILE = index_meta.chunks_file()    # KB_CHUNKS_FILE 可覆盖
SUPPORTED_EXTS = index_meta.SUPPORTED_EXTS
UNSUPPORTED = index_meta.UNSUPPORTED_EXTS

print("加载 embedding 模型 ...", flush=True)
embed_model = SentenceTransformer("BAAI/bge-small-zh-v1.5")
print("模型就绪", flush=True)

# ============ 文档解析（与 app.py 一致） ============
def extract_lines(page, y_tol=3.0):
    def is_data(w):
        return bool(re.match(r'^[\d±—✓\-\.\(\)]', w))
    words = page.get_text("words")
    words.sort(key=lambda w: (round(w[1], 1), w[0]))
    raw_lines, cur, cur_y = [], [], None
    for w in words:
        y = w[1]
        if cur_y is None or abs(y - cur_y) > y_tol:
            if cur:
                raw_lines.append(cur)
            cur, cur_y = [w[4]], y
        else:
            cur.append(w[4])
    if cur:
        raw_lines.append(cur)
    result = []
    for ws in raw_lines:
        dc = sum(1 for w in ws if is_data(w))
        if dc >= 2 and len(ws) - dc >= 1:
            labels = [w for w in ws if not is_data(w)]
            datas = [w for w in ws if is_data(w)]
            result.append(" ".join(labels + datas))
        else:
            result.append(" ".join(ws))
    return result

def load_pdf(path):
    doc = pymupdf.open(str(path))
    lines = []
    for page in doc:
        lines.extend(extract_lines(page))
    return lines

def load_document(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(str(path))
    elif ext in [".txt", ".md"]:
        return Path(path).read_text(encoding="utf-8", errors="replace").splitlines()
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [para.text for para in d.paragraphs if para.text.strip()]
        for table in d.tables:
            for row in table.rows:
                # 合并单元格(gridSpan)在 row.cells 里会连续重复返回同一单元格，
                # 导致工程名等字段一个块里重复 4~8 次、BM25 词频虚高，把真答案挤出 top-k。
                # 这里做相邻去重，保留顺序，去掉合并单元格造成的连续重复。
                cells = []
                prev = None
                for c in row.cells:
                    t = c.text.strip()
                    if t != prev:
                        cells.append(t)
                    prev = t
                if any(cells):
                    lines.append(" | ".join(cells))
        return lines
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[工作表: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
        wb.close()
        return lines
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        lines = []
        for sh in wb.sheets():
            lines.append(f"[工作表: {sh.name}]")
            for r in range(sh.nrows):
                cells = [str(sh.cell_value(r, c)) for c in range(sh.ncols)]
                if any(c.strip() for c in cells):
                    lines.append(" | ".join(cells))
        return lines
    return []

def split_chunks(lines, chunk_size=400):
    chunks, current, cur_len = [], [], 0
    for line in lines:
        if current and cur_len + len(line) > chunk_size:
            chunks.append("\n".join(current))
            current, cur_len = [], 0
        current.append(line)
        cur_len += len(line)
    if current:
        chunks.append("\n".join(current))
    return chunks

# ============ 切分策略（可扩展，默认 single_400） ============
# single_400   单层固定 400 字符（默认，事实型数据最优）
# parent_child 父子文档（预留：答案跨段时才启用，见 chunk_router.py）
CHUNK_STRATEGY = "single_400"

def split_by_strategy(lines, strategy=None):
    strategy = strategy or CHUNK_STRATEGY
    if strategy == "single_400":
        return split_chunks(lines)
    if strategy == "parent_child":
        # 父子文档：返回子块列表。启用需同步改 chunks.json 结构（加 parent 块 + child→parent 映射）
        # 及检索逻辑（召回子块 → 映射回父块喂 LLM）。见 chunk_router.py 的 split_parent_child。
        raise NotImplementedError(
            "父子文档未启用：需改 chunks.json 结构 + 检索映射。"
            "当前事实型数据实验证明父子文档是负优化，不要启用。")
    raise ValueError(f"未知切分策略 {strategy!r}，可选 single_400 / parent_child")

def chunk_strategy_for(path):
    """文档类型路由钩子（预留）。当前全部走 single_400。
    未来数据异构时改成：散文(占比>0.7)→parent_child，表格→single_400。
    判据见 chunk_router.doc_prose_ratio。"""
    return CHUNK_STRATEGY

def build_index(texts):
    vectors = np.array(embed_model.encode(texts)).astype("float32")
    index = faiss.IndexFlatL2(vectors.shape[1])
    index.add(vectors)
    return index

# ============ 扫描（与 app.py / index_meta 共用同一套规则） ============
files, unsupported = index_meta.supported_files(DOCS_DIR)

print(f"扫描到 {len(files)} 个支持文件，{len(unsupported)} 个不支持格式", flush=True)
if unsupported:
    for f in unsupported[:20]:
        print(f"  ⚠ 跳过: {f.relative_to(DOCS_DIR)}", flush=True)

# ============ 逐文件入库 ============
all_chunks, all_sources = [], []
for i, f in enumerate(files):
    rel = str(f.relative_to(DOCS_DIR))
    try:
        lines = load_document(str(f))
        cs = split_by_strategy(lines, chunk_strategy_for(f))
        if not cs:
            print(f"[{i+1}/{len(files)}] ⚠ {rel}: 无内容", flush=True)
            continue
        all_chunks.extend(cs)
        all_sources.extend([rel] * len(cs))
        print(f"[{i+1}/{len(files)}] {rel}: {len(cs)} 块", flush=True)
    except Exception as e:
        print(f"[{i+1}/{len(files)}] ⚠ {rel}: 失败 {type(e).__name__}", flush=True)

print(f"编码 {len(all_chunks)} 块 ...", flush=True)
index = build_index(all_chunks)

# ============ 存盘（先写 .tmp 再原子替换：重建过程中 app 读到的仍是旧索引，不会读到半个文件） ============
tmp_index = INDEX_FILE.with_suffix(".index.tmp")
tmp_chunks = CHUNKS_FILE.with_suffix(".json.tmp")
with open(tmp_index, "wb") as f:
    f.write(faiss.serialize_index(index).tobytes())
data = [{"text": t, "source": s} for t, s in zip(all_chunks, all_sources)]
json.dump(data, open(tmp_chunks, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
os.replace(tmp_index, INDEX_FILE)
os.replace(tmp_chunks, CHUNKS_FILE)

# 记录 docs 目录指纹：app.py 靠它判断"要不要自动重建"
meta = index_meta.save(len(all_chunks), len(files), DOCS_DIR)
print(f"DONE 知识库 {len(all_chunks)} 块，来自 {len(files)} 个文件", flush=True)
print(f"DONE 指纹 {meta['fingerprint'][:12]}... 已写入 index_meta.json", flush=True)
