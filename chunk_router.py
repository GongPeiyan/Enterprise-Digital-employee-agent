# -*- coding: utf-8 -*-
"""
二元路由切分：散文 → 父子文档；表格 → 单层400
判据：文档的"散文文字占比"（段落文字 / 段落+表格文字）

这是骨架版，你理解后可自己改：
  - PROSE_THRESHOLD 阈值（0.7 = 散文占比超过70%算散文）
  - 父子文档的 child/parent 大小（200/800 字符）
"""
from pathlib import Path

PROSE_THRESHOLD = 0.7
CHILD_SIZE, PARENT_SIZE = 200, 800

def doc_prose_ratio(path):
    """返回文档的散文文字占比 0~1（1=纯散文，0=纯表格）"""
    ext = Path(path).suffix.lower()
    if ext == ".xlsx":
        return 0.0                      # 纯表格
    if ext in (".txt", ".md"):
        return 1.0                      # 纯文本
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        p = sum(len(x.text) for x in d.paragraphs)
        t = sum(len(c.text) for tb in d.tables for r in tb.rows for c in r.cells)
        total = p + t
        return p / total if total else 0.5
    if ext == ".pdf":
        return 0.5                      # PDF 暂按中性（后续可加版面分析）
    return 0.5

def split_flat(lines, size=400):
    chunks, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen + len(ln) > size:
            chunks.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: chunks.append("\n".join(cur))
    return chunks

def split_parent_child(lines, child=CHILD_SIZE, parent=PARENT_SIZE):
    parents, cur, clen = [], [], 0
    for ln in lines:
        if cur and clen + len(ln) > parent:
            parents.append("\n".join(cur)); cur, clen = [], 0
        cur.append(ln); clen += len(ln)
    if cur: parents.append("\n".join(cur))
    children, c2p = [], []
    for pi, p in enumerate(parents):
        cc, cclen = [], 0
        for ln in p.split("\n"):
            if cc and cclen + len(ln) > child:
                children.append("\n".join(cc)); c2p.append(pi); cc, cclen = [], 0
            cc.append(ln); cclen += len(ln)
        if cc: children.append("\n".join(cc)); c2p.append(pi)
    return children, c2p, parents

def route(path):
    """二元路由：散文→parent_child，表格→flat"""
    return "parent_child" if doc_prose_ratio(path) > PROSE_THRESHOLD else "flat"

# ============ 冒烟测试 ============
if __name__ == "__main__":
    DOCS = Path(r"E:\数字员工项目\docs")
    tests = [
        DOCS / "天然气" / "XX城市燃气有限公司员工手册.docx",   # 应判散文
        DOCS / "天然气" / "规划部" / "规划部工程资料" / "示例小区" / "示例路（示例小区）",  # 应判表格
        DOCS / "天然气" / "天然气有关" / "开工报审" / "7.1施工组织方案.docx",  # 应判散文
    ]
    for f in tests:
        if not f.exists():
            print(f"不存在: {f.name}"); continue
        r = doc_prose_ratio(f)
        mode = route(f)
        print(f"散文占比 {r:.2f}  →  {mode:12}  {f.name[:50]}")
